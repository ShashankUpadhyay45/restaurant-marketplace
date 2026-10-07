import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(17)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    elif level == 3:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    return h

def build_audit_doc(filepath):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)
        
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("GOURMETFLEET — FULL-STACK AUDIT REPORT")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("Deep Technical & Product Audit: Discovery Flaws, US-Centric Residue, Geolocation Deficits, and Indian Marketplace Upgrade Strategy")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(20)
    r_meta = p_meta.add_run("Audit Scope: Monorepo Codebase, Data Models, Geolocation, Indian Localization & AI | Date: September 2026")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)
    
    add_header_styled(doc, "1. Executive Audit Summary", 1)
    doc.add_paragraph(
        "A rigorous, line-by-line inspection of the existing GourmetFleet codebase was conducted across frontend components, "
        "Express controllers, Mongoose schemas, pricing algorithms, AI services, and database seed scripts. While the underlying "
        "MERN architecture provides strong foundations—specifically with deterministic server-side pricing, atomic reservation locking, "
        "Socket.IO real-time channels, and Three.js visual assets—the platform exhibits severe prototype limitations and geographic "
        "incongruities that prevent it from functioning as a realistic Indian food-tech marketplace."
    )
    
    # Audit Findings Matrix
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdrs = ["System Component", "Current Implementation Status", "Discovered Vulnerability / Flaw", "Severity & Remediation Required"]
    for i, h in enumerate(hdrs):
        c = table.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "1A365D")
        set_cell_margins(c, 100, 100, 80, 80)
        p = c.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    audit_findings = [
        ("Location Management", "Hardcoded / Static", "Navbar hardcodes 'Metropolis Center (3.5 km)'. Zero browser geolocation support; no city/locality picker.", "CRITICAL: Replace with Geolocation API & India city/locality modal."),
        ("Geospatial Database", "Non-Indexed Flat Coordinates", "Restaurant model stores simple {lat, lng} without GeoJSON format or 2dsphere spatial index. Geospatial queries ($near, $geoNear) impossible.", "HIGH: Migrate to GeoJSON Point format [lng, lat] and create 2dsphere index."),
        ("Seed Data & Addresses", "Fictional US Metros", "Restaurants seeded in 'Metropolis, NY, 10001' with US phone numbers (+1555...) and Western addresses.", "CRITICAL: Replace with authentic multi-city Indian dataset (Dehradun, Delhi, Bengaluru, Mumbai)."),
        ("Currency & Pricing", "US Dollar ($)", "All menu items, checkout calculations, GST, driver tips, and analytics hardcoded in USD ($) instead of Indian Rupees (₹).", "CRITICAL: Convert entire platform to INR (₹) with Indian delivery and packaging tiers."),
        ("Search Engine", "Basic Text Match", "Homepage search lacks debouncing, only queries restaurant name/cuisine, and ignores user's active geographic location.", "HIGH: Implement debounced unified search with categorized dropdown (Cities, Localities, Dishes, Venues)."),
        ("Delivery Feasibility", "Simulated / Static", "Platform assumes universal delivery without checking if user coordinates fall within restaurant's actual delivery radius.", "HIGH: Implement Haversine distance calculator and delivery zone boundary checks."),
        ("AI Assistant", "US Metrics Assumptions", "Owner AI copilot cites dollar revenues ($) and Western operational patterns rather than Indian dining trends.", "MEDIUM: Localize prompts, INR figures, and add customer geographic breakdown analytics."),
        ("Admin Governance", "Global List Only", "Admin dashboard lacks city/locality management and city-level operational performance metrics.", "MEDIUM: Add Indian city registry, locality controls, and restaurant counts per city.")
    ]
    
    for r_idx, (comp, status, flaw, rem) in enumerate(audit_findings):
        row = table.add_row()
        for ci, val in enumerate([comp, status, flaw, rem]):
            c = row.cells[ci]
            c.text = val
            bg = "F7FAFC" if r_idx % 2 == 0 else "FFFFFF"
            if "CRITICAL" in rem:
                bg = "FFF5F5"
            elif "HIGH" in rem:
                bg = "FEFCBF"
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 70, 70)
            p = c.paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    add_header_styled(doc, "2. Detailed Analysis of Critical Existing Flaws", 1)
    
    flaws = [
        ("A. Hardcoded Static Location ('Metropolis Center')",
         "In client/src/components/layout/Navbar.tsx, the location display is hardcoded as:\n"
         "  <span className=\"font-medium text-stone-800\">Metropolis Center</span>\n"
         "  <span className=\"text-stone-400 ml-1.5 text-[11px]\">(3.5 km delivery radius)</span>\n"
         "This provides zero user interactivity. A user in Dehradun, Delhi, or Bengaluru cannot detect or select their locality, "
         "nor does changing the location filter or sort nearby restaurants."),
        ("B. Fictional US Restaurant & Address Data",
         "The database seed script (server/src/seed/seed.ts) populated mock data using fictional American locations:\n"
         "  - '14 Rue Montorgueil Arts District, Metropolis, NY, 10003'\n"
         "  - '88 Little Tokyo Lane, Metropolis, NY, 10013'\n"
         "  - '742 Evergreen Terrace, Apt 4B, Metropolis, NY, 10001'\n"
         "  - Phone numbers formatted with US area codes (+15551003001).\n"
         "This violates authentic Indian marketplace standards and contradicts the platform's stated FSSAI and Indian GSTIN compliance."),
        ("C. Absence of MongoDB Geospatial (2dsphere) Indexing",
         "The Restaurant schema (server/src/models/Restaurant.ts) stored coordinates as a non-standard sub-object { lat: Number, lng: Number } "
         "without GeoJSON specifications. Consequently, MongoDB could not execute spatial index lookups using $near or $geoNear. "
         "Distance calculations were completely simulated rather than calculated from genuine geospatial coordinates."),
        ("D. Currency Mismatch ($ vs ₹)",
         "Throughout all client views (Cart drawer, Checkout page, Orders tracker, Restaurant detail, Owner dashboard), prices were rendered "
         "in US Dollars ($) with base rates ($30 duck, $18 escargots) that do not reflect Indian purchasing power or restaurant economics. "
         "The pricing engine must operate in Indian Rupees (₹) with localized delivery charges (₹35 base + ₹12/km) and packaging fees (₹20)."),
        ("E. Lack of Search Debouncing & Categorized Autocomplete",
         "The existing search input in client/src/pages/customer/HomePage.tsx executed searches on raw input changes without debouncing, "
         "triggering redundant component re-renders. Furthermore, search was limited to simple text substring matching on restaurant names, "
         "failing to provide categorized suggestions for Dishes, Cuisines, Cities, and Localities.")
    ]
    
    for f_title, f_desc in flaws:
        add_header_styled(doc, f_title, 2)
        p = doc.add_paragraph()
        r = p.add_run(f_desc)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        
    add_header_styled(doc, "3. Complete Upgrade Strategy & Action Plan", 1)
    doc.add_paragraph(
        "To transform GourmetFleet into an authentic Indian culinary marketplace, the following phased upgrades are being implemented:"
    )
    
    actions = [
        ("1. Centralized Indian Location Engine: ", "Author shared/data/indianLocations.ts containing 35+ major cities across North, West, South, and East India with genuine coordinates, locality boundaries (Connaught Place, Rajpur Road, Koramangala, Bandra West), and pincodes."),
        ("2. Real Browser Geolocation & Permission Handling: ", "Implement navigator.geolocation API integration in useLocationStore with full state machine handling: Permission Granted, Permission Denied, Timeout, and Manual Override."),
        ("3. Backend Reverse-Geocoding & Nearby API: ", "Create GET /api/locations/cities, GET /api/locations/search, POST /api/locations/reverse-geocode, and GET /api/restaurants/nearby powered by Haversine distance calculations and MongoDB GeoJSON 2dsphere indexing."),
        ("4. Authentic Indian Multi-City Seed Data: ", "Replace all US demo data with authentic Indian restaurants across Dehradun, Delhi, Bengaluru, and Mumbai featuring authentic Indian regional cuisines (Awadhi Biryani, Mughlai, North Indian, Indo-Chinese, South Indian) with INR (₹) pricing."),
        ("5. Location-Aware Search & Map Discovery: ", "Upgrade search with 300ms debounce and categorized suggestions (Localities, Cities, Dishes, Venues), plus an interactive 'Explore on Map' view showing real distances and delivery zone eligibility."),
        ("6. Owner Customer Geography & Localized AI: ", "Equip restaurant owners with geographic distribution analytics showing order volume by neighborhood, and localize the Owner AI Copilot to analyze Indian business patterns and INR revenue metrics.")
    ]
    
    for title, desc in actions:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(title)
        r1.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        r2 = p.add_run(desc)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
    doc.save(filepath)
    print(f"Successfully generated {filepath}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "GourmetFleet_India_Audit_Report.docx"
    build_audit_doc(target)
