import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(17)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    elif level == 3:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    return h

def build_plan_doc(filepath):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)
        
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("COMPLETE IMPLEMENTATION PLAN & SPECIFICATIONS")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("End-to-End Roadmap, Feature Requirements, Milestones, File Structure, and Acceptance Criteria")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(20)
    r_meta = p_meta.add_run("Target: GourmetFleet Marketplace  |  Stack: MongoDB, Express, React, Node.js, TypeScript  |  Version: 1.0.0")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)
    
    add_header_styled(doc, "1. Project Phases & Milestones Breakdown", 1)
    
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdrs = ["Phase #", "Milestone / Workstream", "Primary Deliverables", "Dependencies", "Acceptance Criteria"]
    for i, h in enumerate(hdrs):
        c = table.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "1A365D")
        set_cell_margins(c, 100, 100, 80, 80)
        p = c.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    phases_data = [
        ("Phase 1-3", "Research, Strategy & Architecture", "Word Docs, Architecture diagrams, schema blueprints", "None", "Complete competitor benchmarking, SWOT, risk register, and data models approved."),
        ("Phase 4", "Premium UI/UX & 3D Web Graphics", "Three.js Hero, Tailwind theme, responsive layouts", "Vite, React, Three.js", "Smooth 60fps canvas, reduced-motion fallbacks, mobile-first responsiveness."),
        ("Phase 5", "Multi-Role Auth & RBAC", "JWT rotation, bcrypt, Auth pages, role guards", "Mongoose, Express, Zod", "Customer, Owner, Admin roles enforced; verification status check on owners."),
        ("Phase 6-7", "Food Discovery & Ordering Engine", "Fuzzy search, transparent cart, checkout, Socket.IO", "Auth, Mongo schemas", "Item subtotal, delivery fee, GST, platform fee explicitly rendered and validated on server."),
        ("Phase 8", "Table Reservation & Concurrency", "Atomic slot booking, QR ticket pass, floor plans", "Mongoose transactions", "Zero double-booking on simultaneous requests; hold countdown timer."),
        ("Phase 9", "Restaurant Owner Portal & Ops", "Kanban orders, table layout, menu editor, analytics", "Socket.IO, Express", "Real-time audio/visual order alert; instant out-of-stock menu toggles."),
        ("Phase 10", "Autonomous Owner AI Assistant", "Tool-calling engine, sales analysis, prompt defense", "Owner auth, AI service", "Only queries authenticated restaurant's data; human approval required for changes."),
        ("Phase 11-13", "Admin Governance, Loyalty & Alerts", "KYC onboarding, points ledger, email/WhatsApp sim", "Order/Reservation state", "Audit logging, points reversal on refund, idempotency in alerts."),
        ("Phase 14-17", "Security, QA, Packaging & State", "Automated tests, Docker compose, final ZIP archive", "All workstreams", "Passing test suite, clean packaging without node_modules, reproducible build.")
    ]
    
    for r_idx, p_row in enumerate(phases_data):
        row = table.add_row()
        for ci, val in enumerate(p_row):
            c = row.cells[ci]
            c.text = val
            bg = "F7FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 70, 70)
            p = c.paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    add_header_styled(doc, "2. Detailed Feature Specifications", 1)
    
    features = [
        ("2.1 Customer Ordering Journey", 
         "• Multi-factor Discovery: Geolocation filtering, cuisine tags, dietary toggles (Vegetarian, Vegan, Gluten-Free, Halal, Chef Special).\n"
         "• Transparent Item Customization: Selection of size variants (e.g. Regular vs Large) and nested add-ons with individual pricing.\n"
         "• Dynamic Fee Calculator: Calculates distance-based delivery tier ($2.00 base + $0.80/km), restaurant packaging, 5% food GST, 18% services GST, and $1.50 platform fee.\n"
         "• Live Order Tracking: Real-time Socket.IO synchronization through 6 distinct order phases with simulated realistic delivery steps."),
        ("2.2 Real-time Table Reservation System",
         "• Inventory-Driven Availability: Tables configured with seating capacity (2-top, 4-top, 6-top, VIP Booth, Terrace).\n"
         "• Atomic Slot Locking: Server evaluates table availability at exact booking timestamp + 90-minute dining duration. Atomic locking prevents split-second race conditions.\n"
         "• Guest Pass Generation: Generates scannable SVG/Canvas QR ticket containing booking token, party size, table number, and arrival window.\n"
         "• No-Show Mitigation: 15-minute grace window timer, calendar sync (.ics), and automated arrival reminders."),
        ("2.3 Restaurant Owner Operating System",
         "• Live Kitchen Display (KDS) & Order Kanban: Columnar interface (Incoming, Accepted/Cooking, Ready for Dispatch, In Transit).\n"
         "• Real-time Floor Plan View: Live visual status of dining room tables (Available, Reserved, Seated, Dirty/Cleaning).\n"
         "• Menu Stock Controls: Instant toggle for ingredient outages to immediately disable affected items in customer cart.\n"
         "• Financial Analytics: Itemized payouts subtracting transparent platform commission (e.g. 10%) and net earnings export."),
        ("2.4 Owner AI Business Assistant",
         "• Tool-Calling Architecture: AI model provided with deterministic schemas to execute read queries against sales, occupancy, and top dishes.\n"
         "• Strict Multi-Tenant Isolation: Every tool call automatically binds `restaurantId: req.user.restaurantId`; cannot leak cross-restaurant data.\n"
         "• Human-in-the-Loop Execution: Any operational action (e.g. applying a 15% discount for slow Wednesday lunches) outputs an executable proposal requiring explicit owner confirmation."),
        ("2.5 Platform Admin Governance",
         "• Merchant Verification Queue: Review restaurant FSSAI license, tax ID (GSTIN), and hygiene certification before enabling live listings.\n"
         "• Commission Rate Customizer: Capability to assign standard default commission (10%) or negotiated custom rates for anchor restaurant partners.\n"
         "• Audit Trail: Every critical administrative and moderation action logged with timestamp, admin ID, IP, and reason.")
    ]
    
    for f_title, f_desc in features:
        add_header_styled(doc, f_title, 2)
        doc.add_paragraph(f_desc)
        
    add_header_styled(doc, "3. Directory & File Organization Structure", 1)
    doc.add_paragraph(
        "The project is structured as a clean, cohesive monorepo ensuring complete separation of concerns while sharing "
        "type contracts between frontend and backend:"
    )
    
    file_tree = (
        "restaurant-marketplace/\n"
        "├── client/                           # Frontend React + TypeScript + Vite Application\n"
        "│   ├── src/\n"
        "│   │   ├── components/               # 3D Canvas, Customer UI, Owner Dashboard, Admin UI\n"
        "│   │   ├── pages/                    # Home, RestaurantDetail, Checkout, Orders, Bookings\n"
        "│   │   ├── store/                    # Zustand state management (auth, cart, notifications)\n"
        "│   │   ├── services/                 # API calls, Socket.IO client, QR pass generator\n"
        "│   │   └── types/                    # Frontend type models\n"
        "│   ├── index.html\n"
        "│   ├── package.json\n"
        "│   ├── tailwind.config.js\n"
        "│   └── vite.config.ts\n"
        "├── server/                           # Backend Node.js + Express + TypeScript Application\n"
        "│   ├── src/\n"
        "│   │   ├── config/                   # MongoDB connection, Redis, environment validation\n"
        "│   │   ├── controllers/              # Auth, Restaurant, Order, Reservation, AI, Admin\n"
        "│   │   ├── middleware/               # requireAuth, requireRole, rateLimit, errorHandler\n"
        "│   │   ├── models/                   # Mongoose Schemas (User, Restaurant, Order, etc.)\n"
        "│   │   ├── services/                 # PricingEngine, ReservationEngine, AIService\n"
        "│   │   ├── sockets/                  # Socket.IO connection and room dispatchers\n"
        "│   │   └── tests/                    # Vitest / Jest automated test suites\n"
        "│   ├── seed/                         # Rich seed script with realistic sample restaurants\n"
        "│   ├── package.json\n"
        "│   └── tsconfig.json\n"
        "├── shared/                           # Shared TypeScript contracts and Zod schemas\n"
        "├── docker-compose.yml                # Multi-container orchestration (Mongo, Redis, API, Web)\n"
        "└── README.md                         # Comprehensive deployment & operation instructions"
    )
    
    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(8)
    r_c = p_code.add_run(file_tree)
    r_c.font.name = "Consolas"
    r_c.font.size = Pt(8.5)
    r_c.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    add_header_styled(doc, "4. Rigorous Acceptance Criteria & Quality Gates", 1)
    
    criteria_table = doc.add_table(rows=1, cols=4)
    criteria_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdrs = ["Category", "Requirement & Target", "Validation Method", "Pass / Fail Gate"]
    for i, h in enumerate(c_hdrs):
        c = criteria_table.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "2B6CB0")
        set_cell_margins(c, 100, 100, 80, 80)
        p = c.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    crit_rows = [
        ("Concurrency", "Zero double bookings under 10 concurrent requests for same table slot", "Automated concurrency test simulating simultaneous booking POSTs", "Strict 1 success, 9 slot-unavailable failures"),
        ("Price Integrity", "Server recalculates all items, addons, taxes; rejects client-altered totals", "Malicious POST payload with $0.01 price attempt", "Server overrides or rejects with 400 Bad Request"),
        ("Tenant Isolation", "Owner A cannot read or mutate Owner B's orders or analytics", "Automated API test invoking Owner A's token on Owner B's endpoints", "HTTP 403 Forbidden with security audit log entry"),
        ("Data Consistency", "Cancelled/refunded orders must reverse awarded loyalty points", "Order refund lifecycle test inspecting loyalty ledger balance", "Ledger records matching negative offset entry; no negative balance"),
        ("Responsive UX", "Seamless execution on desktop (1920x1080), tablet (768x1024), and mobile (375x812)", "Lighthouse & Playwright viewport responsive tests", "Zero horizontal scrollbars, responsive sticky cart & touch-friendly tap targets")
    ]
    
    for r_idx, (cat, req, val, gate) in enumerate(crit_rows):
        row = criteria_table.add_row()
        for ci, val_text in enumerate([cat, req, val, gate]):
            cell = row.cells[ci]
            cell.text = val_text
            bg = "F7FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(cell, bg)
            set_cell_margins(cell, 70, 70, 70, 70)
            p = cell.paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            
    doc.save(filepath)
    print(f"Successfully generated {filepath}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "Complete_Implementation_Plan.docx"
    build_plan_doc(target)
