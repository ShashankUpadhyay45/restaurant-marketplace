import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(6)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(17)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    elif level == 2:
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    elif level == 3:
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        for run in h.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    return h

def build_plan_doc(filepath):
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Inches(1)
        sec.bottom_margin = Inches(1)
        sec.left_margin = Inches(1)
        sec.right_margin = Inches(1)
        
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("GOURMETFLEET — INDIA LOCALIZATION & IMPLEMENTATION PLAN")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
    
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("Technical Roadmap, Geolocation Architecture, MongoDB Geospatial Schemas, Unified Search, and Indian Marketplace Upgrade")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(20)
    r_meta = p_meta.add_run("Author: Principal MERN Architect & Geolocation Engineer  |  Target: GourmetFleet India 2.0  |  Version: 2.0.0")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = RGBColor(0xA0, 0xAE, 0xC0)
    
    add_header_styled(doc, "1. Upgrade Objectives & Architectural Scope", 1)
    doc.add_paragraph(
        "This implementation plan defines the complete upgrade path required to transform GourmetFleet from a general prototype "
        "into a location-aware, multi-city Indian culinary marketplace. The upgraded system seamlessly integrates browser geolocation, "
        "manual city/locality selection, MongoDB 2dsphere geospatial indexing, authentic Indian culinary datasets, INR (₹) financial models, "
        "interactive map exploration, and localized restaurant owner intelligence."
    )
    
    # Roadmap Table
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdrs = ["Milestone", "Core Deliverables", "Architecture / Technology Stack", "Target Verification Gate"]
    for i, h in enumerate(hdrs):
        c = table.rows[0].cells[i]
        c.text = h
        set_cell_background(c, "1A365D")
        set_cell_margins(c, 100, 100, 80, 80)
        p = c.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    milestones = [
        ("M1: India Location Dataset", "35+ Indian cities, authentic localities, pincodes & coordinates", "shared/data/indianLocations.ts", "Accurate GPS coordinates matching real Indian city centers."),
        ("M2: Geospatial Database", "GeoJSON Point schema migration & 2dsphere indexing", "MongoDB, Mongoose 8.9", "Indexed spatial lookups supporting $near and $geoNear queries."),
        ("M3: Geolocation State & UX", "Browser Geolocation API, permission handling, location modal", "Zustand, React 18, HTML5 Geolocation", "Permission state machine: Granted, Denied, Timeout, Manual Override."),
        ("M4: Location-Aware APIs", "Reverse-geocoding, /api/locations/search, /api/restaurants/nearby", "Express, Node.js 22, Haversine Engine", "Sub-20ms spatial queries returning distance and delivery eligibility."),
        ("M5: Multi-City Indian Seed", "Realistic Indian restaurants (Dehradun, Delhi, Bengaluru, Mumbai)", "Seed Script, Regional Menus, +91 Phones", "Demonstration data clearly labeled with authentic regional cuisines."),
        ("M6: INR Pricing & Search", "Unified search bar with debounce (300ms), INR (₹) price engine", "Tailwind CSS, React Hook Form, PricingEngine", "Categorized autocomplete (Venues, Dishes, Cuisines, Localities)."),
        ("M7: Map View & AI Analytics", "Interactive map view, customer geography distribution, AI localization", "Leaflet / SVG Radar, Owner AI Service", "Visual map markers and owner customer distribution breakdown."),
        ("M8: QA & Packaging", "Automated test suites, build validation, ZIP packaging", "Vitest, tsc, vite build, Python Packager", "100% test pass rate and clean downloadable ZIP archive.")
    ]
    
    for r_idx, (m_id, deliv, stack, gate) in enumerate(milestones):
        row = table.add_row()
        for ci, val in enumerate([m_id, deliv, stack, gate]):
            c = row.cells[ci]
            c.text = val
            bg = "F7FAFC" if r_idx % 2 == 0 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 70, 70)
            p = c.paragraphs[0]
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(8)
            p.runs[0].font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)
    
    add_header_styled(doc, "2. Comprehensive Indian City & Locality Hierarchy", 1)
    doc.add_paragraph(
        "The platform establishes a centralized location authority (shared/data/indianLocations.ts) covering four primary regions of India, "
        "with authentic geographic coordinates, representative localities, and postal codes:"
    )
    
    cities_detail = [
        ("North India: ", "Dehradun (30.3165°N, 78.0322°E - Localities: Rajpur Road, Jakhan, Ballupur, Clement Town, Prem Nagar, Sahastradhara Road), Delhi NCR (28.6139°N, 77.2090°E - Connaught Place, Saket, Hauz Khas, Vasant Kunj, Lajpat Nagar, Karol Bagh, Dwarka), Gurugram (DLF Cyber City, Sector 29, Golf Course Road), Noida (Sector 18, Sector 62), Chandigarh (Sector 17, Sector 35), Jaipur (C-Scheme, Malviya Nagar, Mansarovar), Lucknow (Hazratganj, Gomti Nagar, Aliganj), Amritsar, Ludhiana, Agra, Varanasi, Jammu."),
        ("West India: ", "Mumbai (19.0760°N, 72.8777°E - Localities: Bandra West, Andheri West, Powai, Lower Parel, Colaba, Juhu, Thane), Pune (18.5204°N, 73.8567°E - Koregaon Park, Baner, Kothrud, Viman Nagar, Hinjawadi), Ahmedabad (SG Highway, Bodakdev, Satellite), Surat, Vadodara, Nashik, Nagpur, Indore, Bhopal, Goa/Panaji (Miramar, Fontainhas, Candolim)."),
        ("South India: ", "Bengaluru (12.9716°N, 77.5946°E - Localities: Koramangala, Indiranagar, HSR Layout, Whitefield, MG Road, Jayanagar, Electronic City), Hyderabad (17.3850°N, 78.4867°E - Banjara Hills, Jubilee Hills, Gachibowli, Hitec City, Madhapur), Chennai (T. Nagar, Nungambakkam, Adyar, Velachery), Kochi (Marine Drive, Panampilly Nagar, Fort Kochi), Thiruvananthapuram, Coimbatore, Mysuru, Mangaluru, Visakhapatnam."),
        ("East India: ", "Kolkata (22.5726°N, 88.3639°E - Localities: Park Street, Salt Lake, Ballygunge, New Town, Alipore), Bhubaneswar (Saheed Nagar, Patia), Patna (Kankarbagh, Boring Road), Ranchi, Guwahati (GS Road, Christian Basti).")
    ]
    
    for r_name, r_content in cities_detail:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(r_name)
        r1.bold = True
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        r2 = p.add_run(r_content)
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        
    add_header_styled(doc, "3. Database Geospatial Schema (GeoJSON & 2dsphere)", 1)
    doc.add_paragraph(
        "To ensure precise geospatial calculations, the MongoDB Restaurant model is updated to adhere strictly to GeoJSON Point format "
        "[longitude, latitude], indexed with a 2dsphere index:"
    )
    
    schema_code = (
        "// Updated Restaurant Mongoose Schema\n"
        "const RestaurantSchema = new Schema<IRestaurant>({\n"
        "  name: { type: String, required: true, trim: true, index: true },\n"
        "  slug: { type: String, required: true, unique: true, lowercase: true },\n"
        "  address: {\n"
        "    addressLine1: { type: String, required: true },\n"
        "    addressLine2: String,\n"
        "    locality: { type: String, required: true, index: true },\n"
        "    city: { type: String, required: true, index: true },\n"
        "    state: { type: String, required: true },\n"
        "    postalCode: { type: String, required: true, index: true },\n"
        "    country: { type: String, default: 'India' }\n"
        "  },\n"
        "  location: {\n"
        "    type: { type: String, enum: ['Point'], default: 'Point' },\n"
        "    coordinates: { type: [Number], required: true } // [longitude, latitude]\n"
        "  },\n"
        "  maxDeliveryRadiusKm: { type: Number, default: 8.0 },\n"
        "  commissionRate: { type: Number, default: 10.0 }\n"
        "});\n\n"
        "// Geospatial 2dsphere index for ultra-fast spatial search\n"
        "RestaurantSchema.index({ location: '2dsphere' });\n"
        "RestaurantSchema.index({ 'address.city': 1, 'address.locality': 1 });"
    )
    
    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_before = Pt(4)
    p_code.paragraph_format.space_after = Pt(8)
    rc = p_code.add_run(schema_code)
    rc.font.name = "Consolas"
    rc.font.size = Pt(8)
    rc.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    add_header_styled(doc, "4. Geolocation Priority & Permission State Machine", 1)
    doc.add_paragraph(
        "The frontend location store (client/src/store/useLocationStore.ts) implements the strict resolution priority:\n"
        "  1. Explicit User Selection (Manual city/locality selection overrides all defaults)\n"
        "  2. Current Detected Location (GPS coordinates reverse-geocoded to closest Indian city)\n"
        "  3. Previously Persisted Preference (Restored from localStorage upon return)\n"
        "  4. Default Indian City (Dehradun, Uttarakhand or Delhi NCR as safe baseline)\n\n"
        "Permission states are surfaced with clear user messaging:\n"
        "• Permission Granted: 'Using your live location: Rajpur Road, Dehradun'\n"
        "• Permission Denied: 'Location access disabled. Search for your city or area manually.'\n"
        "• Timeout / Unavailable: Fast retry trigger and instant manual city selector fallback."
    )
    
    doc.save(filepath)
    print(f"Successfully generated {filepath}")

if __name__ == "__main__":
    import sys
    target = sys.argv[1] if len(sys.argv) > 1 else "GourmetFleet_India_Implementation_Plan.docx"
    build_plan_doc(target)
